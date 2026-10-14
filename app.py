import io
import os
import secrets
import sqlite3
from datetime import datetime
from decimal import ROUND_HALF_UP, Decimal

import fitz
from flask import Flask, jsonify, request, send_file, send_from_directory, session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from werkzeug.security import check_password_hash, generate_password_hash

from auth import admin_required, login_required
from db import close_connection, get_db, init_db

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY') or secrets.token_hex(32)
app.json.ensure_ascii = False

MONEY = Decimal('0.01')
QTY = Decimal('0.001')
PRICE = Decimal('0.0001')
STATUSES = ('unpaid', 'paid', 'cancelled')
STATUS_LABELS = {
    'unpaid': 'Неплатена',
    'paid': 'Платена',
    'cancelled': 'Откажана',
}
HEADER_FILL = PatternFill('solid', fgColor='1F4D3A')
HEADER_FONT = Font(color='FFFFFF', bold=True)


@app.teardown_appcontext
def _close_connection(exception):
    close_connection(exception)


def _ready():
    if not getattr(app, '_db_ready', False):
        init_db()
        app._db_ready = True


@app.before_request
def _prepare():
    _ready()


def money(value):
    return Decimal(str(value)).quantize(MONEY, rounding=ROUND_HALF_UP)


def denar(value):
    return money(value).quantize(Decimal(1), rounding=ROUND_HALF_UP)


def decimal_or_none(value, places):
    if value is None or value == '':
        return None
    text = str(value).strip().replace(' ', '').replace('ден', '')
    if ',' in text and '.' in text:
        text = text.replace('.', '').replace(',', '.')
    else:
        text = text.replace(',', '.')
    return Decimal(text).quantize(places, rounding=ROUND_HALF_UP)


def parse_number(value, places=MONEY, required=False, label='Износ'):
    if value is None or value == '':
        if required:
            raise ValueError(f'{label} е задолжителен')
        return None
    try:
        return decimal_or_none(value, places)
    except Exception as exc:
        raise ValueError(f'{label} не е број') from exc


def parse_date(value, required=False, label='Датум'):
    text = (value or '').strip()
    if not text:
        if required:
            raise ValueError(f'{label} е задолжителен')
        return None
    for fmt in ('%Y-%m-%d', '%d.%m.%Y', '%d/%m/%Y'):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    raise ValueError(f'{label} треба да биде ДД.ММ.ГГГГ')


def optional(value):
    text = (value or '').strip() if isinstance(value, str) else value
    return text or None


def float_or_none(value):
    return None if value is None else float(value)


def display_date(value):
    if not value:
        return ''
    try:
        return datetime.strptime(value, '%Y-%m-%d').strftime('%d.%m.%Y')
    except ValueError:
        return value


def row_dict(row):
    return dict(row) if row is not None else None


def bad(message, code=400):
    return jsonify({'error': message}), code


@app.route('/')
def index():
    return send_from_directory('static', 'index.html')


@app.route('/api/login', methods=['POST'])
def login():
    data = request.get_json(silent=True) or {}
    db = get_db()
    user = db.execute(
        'SELECT * FROM users WHERE username = ?',
        ((data.get('username') or '').strip(),),
    ).fetchone()
    if not user or not check_password_hash(user['password_hash'], data.get('password') or ''):
        return bad('Погрешно корисничко име или лозинка', 401)
    session['user_id'] = user['id']
    session['username'] = user['username']
    session['role'] = user['role']
    return jsonify({'id': user['id'], 'username': user['username'], 'role': user['role']})


@app.route('/api/logout', methods=['POST'])
def logout():
    session.clear()
    return jsonify({'ok': True})


@app.route('/api/me')
def me():
    if 'user_id' not in session:
        return bad('Не сте најавени', 401)
    return jsonify({
        'id': session['user_id'],
        'username': session['username'],
        'role': session['role'],
    })


@app.route('/api/me/password', methods=['PUT'])
@login_required
def change_own_password():
    data = request.get_json(silent=True) or {}
    new_password = (data.get('password') or '').strip()
    if not new_password:
        return bad('Лозинката е задолжителна')
    db = get_db()
    user = db.execute('SELECT * FROM users WHERE id = ?', (session['user_id'],)).fetchone()
    if not user or not check_password_hash(user['password_hash'], data.get('current_password') or ''):
        return bad('Сегашната лозинка не е точна')
    db.execute(
        'UPDATE users SET password_hash = ? WHERE id = ?',
        (generate_password_hash(new_password), session['user_id']),
    )
    db.commit()
    return jsonify({'ok': True})


@app.route('/api/users', methods=['GET'])
@admin_required
def get_users():
    rows = get_db().execute(
        'SELECT id, username, role, created_at FROM users ORDER BY id'
    ).fetchall()
    return jsonify([dict(row) for row in rows])


@app.route('/api/users', methods=['POST'])
@admin_required
def create_user():
    data = request.get_json(silent=True) or {}
    username = (data.get('username') or '').strip()
    password = (data.get('password') or '').strip()
    role = data.get('role') or 'user'
    if not username or not password:
        return bad('Корисничко име и лозинка се задолжителни')
    if role not in ('admin', 'user'):
        return bad('Неважечка улога')
    db = get_db()
    try:
        cursor = db.execute(
            'INSERT INTO users (username, password_hash, role) VALUES (?, ?, ?)',
            (username, generate_password_hash(password), role),
        )
        db.commit()
    except sqlite3.IntegrityError:
        return bad('Корисничкото име веќе постои', 409)
    row = db.execute(
        'SELECT id, username, role, created_at FROM users WHERE id = ?',
        (cursor.lastrowid,),
    ).fetchone()
    return jsonify(dict(row)), 201


@app.route('/api/users/<int:user_id>', methods=['PUT'])
@admin_required
def update_user(user_id):
    data = request.get_json(silent=True) or {}
    db = get_db()
    user = db.execute('SELECT * FROM users WHERE id = ?', (user_id,)).fetchone()
    if not user:
        return bad('Корисникот не постои', 404)
    username = (data.get('username') or user['username']).strip()
    role = data.get('role') or user['role']
    if not username:
        return bad('Корисничкото име е задолжително')
    if role not in ('admin', 'user'):
        return bad('Неважечка улога')
    if user_id == session['user_id'] and role != 'admin':
        return bad('Не можете да си ја одземете администраторската улога')
    password = (data.get('password') or '').strip()
    try:
        if password:
            db.execute(
                'UPDATE users SET username = ?, role = ?, password_hash = ? WHERE id = ?',
                (username, role, generate_password_hash(password), user_id),
            )
        else:
            db.execute(
                'UPDATE users SET username = ?, role = ? WHERE id = ?',
                (username, role, user_id),
            )
        db.commit()
    except sqlite3.IntegrityError:
        return bad('Корисничкото име веќе постои', 409)
    if user_id == session['user_id']:
        session['username'] = username
        session['role'] = role
    row = db.execute(
        'SELECT id, username, role, created_at FROM users WHERE id = ?',
        (user_id,),
    ).fetchone()
    return jsonify(dict(row))


@app.route('/api/users/<int:user_id>', methods=['DELETE'])
@admin_required
def delete_user(user_id):
    if user_id == session['user_id']:
        return bad('Не можете да се избришете себеси')
    db = get_db()
    deleted = db.execute('DELETE FROM users WHERE id = ?', (user_id,)).rowcount
    db.commit()
    if not deleted:
        return bad('Корисникот не постои', 404)
    return jsonify({'ok': True})


@app.route('/api/tax-types', methods=['GET'])
@login_required
def get_tax_types():
    rows = get_db().execute('SELECT * FROM tax_types ORDER BY rate, name').fetchall()
    return jsonify([dict(row) for row in rows])


@app.route('/api/tax-types', methods=['POST'])
@login_required
def create_tax_type():
    data = request.get_json(silent=True) or {}
    name = (data.get('name') or '').strip()
    if not name:
        return bad('Називот е задолжителен')
    try:
        rate = parse_number(data.get('rate'), MONEY, required=True, label='Стапка')
    except ValueError as exc:
        return bad(str(exc))
    if rate is None or rate < 0 or rate > 100:
        return bad('Стапката треба да биде помеѓу 0 и 100')
    stored = (rate / Decimal(100)).quantize(Decimal('0.0001'), rounding=ROUND_HALF_UP)
    db = get_db()
    try:
        cursor = db.execute(
            'INSERT INTO tax_types (name, rate) VALUES (?, ?)',
            (name, float(stored)),
        )
        db.commit()
    except sqlite3.IntegrityError:
        return bad('Даночна стапка со тоа име веќе постои', 409)
    return jsonify(row_dict(db.execute(
        'SELECT * FROM tax_types WHERE id = ?', (cursor.lastrowid,)
    ).fetchone())), 201


@app.route('/api/tax-types/<int:tax_id>', methods=['PUT'])
@login_required
def update_tax_type(tax_id):
    data = request.get_json(silent=True) or {}
    db = get_db()
    current = db.execute('SELECT * FROM tax_types WHERE id = ?', (tax_id,)).fetchone()
    if not current:
        return bad('Даночната стапка не постои', 404)
    name = (data.get('name') or current['name']).strip()
    if not name:
        return bad('Називот е задолжителен')
    try:
        if 'rate' in data:
            rate = parse_number(data.get('rate'), MONEY, required=True, label='Стапка')
            if rate is None or rate < 0 or rate > 100:
                return bad('Стапката треба да биде помеѓу 0 и 100')
            stored = float((rate / Decimal(100)).quantize(Decimal('0.0001'), rounding=ROUND_HALF_UP))
        else:
            stored = current['rate']
        db.execute(
            'UPDATE tax_types SET name = ?, rate = ? WHERE id = ?',
            (name, stored, tax_id),
        )
        db.commit()
    except ValueError as exc:
        return bad(str(exc))
    except sqlite3.IntegrityError:
        return bad('Даночна стапка со тоа име веќе постои', 409)
    return jsonify(row_dict(db.execute(
        'SELECT * FROM tax_types WHERE id = ?', (tax_id,)
    ).fetchone()))


@app.route('/api/tax-types/<int:tax_id>', methods=['DELETE'])
@login_required
def delete_tax_type(tax_id):
    db = get_db()
    used = db.execute(
        'SELECT COUNT(*) FROM invoice_lines WHERE tax_type_id = ?', (tax_id,)
    ).fetchone()[0]
    if used:
        return bad('Стапката се користи на фактури и не може да се избрише')
    deleted = db.execute('DELETE FROM tax_types WHERE id = ?', (tax_id,)).rowcount
    db.commit()
    if not deleted:
        return bad('Даночната стапка не постои', 404)
    return jsonify({'ok': True})


def partner_accounts(db, partner_id):
    return [dict(row) for row in db.execute(
        'SELECT id, bank, account_number FROM bank_accounts WHERE partner_id = ? ORDER BY id',
        (partner_id,),
    )]


def save_accounts(db, partner_id, accounts):
    db.execute('DELETE FROM bank_accounts WHERE partner_id = ?', (partner_id,))
    rows = []
    for account in accounts or []:
        number = (account.get('account_number') or '').strip()
        if not number:
            continue
        rows.append((partner_id, optional(account.get('bank')), number))
    if rows:
        db.executemany(
            'INSERT INTO bank_accounts (partner_id, bank, account_number) VALUES (?, ?, ?)',
            rows,
        )


def partner_payload(row):
    item = dict(row)
    db = get_db()
    item['place_count'] = db.execute(
        'SELECT COUNT(*) FROM places WHERE partner_id = ?', (row['id'],)
    ).fetchone()[0]
    item['accounts'] = partner_accounts(db, row['id'])
    return item


@app.route('/api/partners', methods=['GET'])
@login_required
def get_partners():
    rows = get_db().execute('SELECT * FROM partners ORDER BY name').fetchall()
    return jsonify([partner_payload(row) for row in rows])


def read_partner(data, current=None):
    name = (data.get('name') if 'name' in data else (current['name'] if current else '')).strip()
    if not name:
        raise ValueError('Називот на партнерот е задолжителен')

    def field(key):
        if key in data:
            return optional(data.get(key))
        return current[key] if current else None

    return {
        'name': name,
        'edb': field('edb'),
        'embs': field('embs'),
        'address': field('address'),
        'city': field('city'),
        'country': field('country'),
        'phone': field('phone'),
        'mobile': field('mobile'),
        'email': field('email'),
        'notes': field('notes'),
    }


@app.route('/api/partners', methods=['POST'])
@login_required
def create_partner():
    data = request.get_json(silent=True) or {}
    try:
        parsed = read_partner(data)
    except ValueError as exc:
        return bad(str(exc))
    db = get_db()
    cursor = db.execute(
        '''INSERT INTO partners
           (name, edb, embs, address, city, country, phone, mobile, email, notes)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)''',
        tuple(parsed.values()),
    )
    save_accounts(db, cursor.lastrowid, data.get('accounts'))
    db.commit()
    return jsonify(partner_payload(db.execute(
        'SELECT * FROM partners WHERE id = ?', (cursor.lastrowid,)
    ).fetchone())), 201


@app.route('/api/partners/<int:partner_id>', methods=['PUT'])
@login_required
def update_partner(partner_id):
    data = request.get_json(silent=True) or {}
    db = get_db()
    current = db.execute('SELECT * FROM partners WHERE id = ?', (partner_id,)).fetchone()
    if not current:
        return bad('Партнерот не постои', 404)
    try:
        parsed = read_partner(data, current)
    except ValueError as exc:
        return bad(str(exc))
    db.execute(
        '''UPDATE partners SET name=?, edb=?, embs=?, address=?, city=?, country=?,
           phone=?, mobile=?, email=?, notes=? WHERE id=?''',
        (*parsed.values(), partner_id),
    )
    if 'accounts' in data:
        save_accounts(db, partner_id, data.get('accounts'))
    db.commit()
    return jsonify(partner_payload(db.execute(
        'SELECT * FROM partners WHERE id = ?', (partner_id,)
    ).fetchone()))


@app.route('/api/partners/<int:partner_id>', methods=['DELETE'])
@login_required
def delete_partner(partner_id):
    db = get_db()
    invoices = db.execute(
        'SELECT COUNT(*) FROM invoices WHERE partner_id = ?', (partner_id,)
    ).fetchone()[0]
    if invoices:
        return bad('Партнерот има фактури и не може да се избрише')
    db.execute('DELETE FROM places WHERE partner_id = ?', (partner_id,))
    deleted = db.execute('DELETE FROM partners WHERE id = ?', (partner_id,)).rowcount
    db.commit()
    if not deleted:
        return bad('Партнерот не постои', 404)
    return jsonify({'ok': True})


def place_payload(row):
    item = dict(row)
    item['partner_name'] = get_db().execute(
        'SELECT name FROM partners WHERE id = ?', (row['partner_id'],)
    ).fetchone()['name']
    return item


@app.route('/api/articles', methods=['GET'])
@login_required
def get_articles():
    partner_id = request.args.get('partner_id')
    sql = '''SELECT a.*, p.name AS partner_name, t.name AS tax_type_name, t.rate AS tax_rate
             FROM articles a
             JOIN partners p ON p.id = a.partner_id
             LEFT JOIN tax_types t ON t.id = a.tax_type_id'''
    params = []
    if partner_id:
        sql += ' WHERE a.partner_id = ?'
        params.append(partner_id)
    sql += ' ORDER BY a.code, a.name'
    return jsonify([dict(row) for row in get_db().execute(sql, params)])


def read_article(data, current=None):
    partner_id = data.get('partner_id', current['partner_id'] if current else None)
    if not partner_id:
        raise ValueError('Партнерот е задолжителен')
    name = text_value(data, 'name', current)
    if not name:
        raise ValueError('Називот на ставката е задолжителен')
    code = optional(data.get('code')) if 'code' in data else (current['code'] if current else None)
    unit = optional(data.get('unit')) if 'unit' in data else (current['unit'] if current else None)
    if 'tax_type_id' in data:
        tax_type_id = data.get('tax_type_id') or None
    else:
        tax_type_id = current['tax_type_id'] if current else None
    if not tax_type_id:
        raise ValueError('Даночната стапка е задолжителна')
    return {
        'partner_id': int(partner_id),
        'code': code,
        'name': name,
        'unit': unit,
        'tax_type_id': int(tax_type_id),
    }


@app.route('/api/articles', methods=['POST'])
@login_required
def create_article():
    data = request.get_json(silent=True) or {}
    db = get_db()
    try:
        parsed = read_article(data)
    except (ValueError, TypeError) as exc:
        return bad(str(exc) or 'Невалидна ставка')
    if not db.execute('SELECT id FROM partners WHERE id = ?', (parsed['partner_id'],)).fetchone():
        return bad('Партнерот не постои', 404)
    if not db.execute('SELECT id FROM tax_types WHERE id = ?', (parsed['tax_type_id'],)).fetchone():
        return bad('Даночната стапка не постои', 404)
    cursor = db.execute(
        'INSERT INTO articles (partner_id, code, name, unit, tax_type_id) VALUES (?, ?, ?, ?, ?)',
        tuple(parsed.values()),
    )
    db.commit()
    return jsonify(dict(db.execute('SELECT * FROM articles WHERE id = ?', (cursor.lastrowid,)).fetchone())), 201


@app.route('/api/articles/<int:article_id>', methods=['PUT'])
@login_required
def update_article(article_id):
    data = request.get_json(silent=True) or {}
    db = get_db()
    current = db.execute('SELECT * FROM articles WHERE id = ?', (article_id,)).fetchone()
    if not current:
        return bad('Ставката не постои', 404)
    try:
        parsed = read_article(data, current)
        if not db.execute('SELECT id FROM tax_types WHERE id = ?', (parsed['tax_type_id'],)).fetchone():
            return bad('Даночната стапка не постои', 404)
        db.execute(
            'UPDATE articles SET partner_id=?, code=?, name=?, unit=?, tax_type_id=? WHERE id=?',
            (*parsed.values(), article_id),
        )
        db.commit()
    except (ValueError, TypeError) as exc:
        return bad(str(exc) or 'Невалидна ставка')
    return jsonify(dict(db.execute('SELECT * FROM articles WHERE id = ?', (article_id,)).fetchone()))


@app.route('/api/articles/<int:article_id>', methods=['DELETE'])
@login_required
def delete_article(article_id):
    db = get_db()
    db.execute('UPDATE invoice_lines SET article_id = NULL WHERE article_id = ?', (article_id,))
    deleted = db.execute('DELETE FROM articles WHERE id = ?', (article_id,)).rowcount
    db.commit()
    if not deleted:
        return bad('Ставката не постои', 404)
    return jsonify({'ok': True})


@app.route('/api/places', methods=['GET'])
@login_required
def get_places():
    partner_id = request.args.get('partner_id')
    sql = 'SELECT * FROM places'
    params = []
    if partner_id:
        sql += ' WHERE partner_id = ?'
        params.append(partner_id)
    sql += ' ORDER BY customer_number'
    rows = get_db().execute(sql, params).fetchall()
    return jsonify([place_payload(row) for row in rows])


def text_value(data, key, current):
    if key in data:
        return (data.get(key) or '').strip()
    return (current[key] or '').strip() if current else ''


def read_place(data, current=None):
    partner_id = data.get('partner_id', current['partner_id'] if current else None)
    if not partner_id:
        raise ValueError('Партнерот е задолжителен')
    name = text_value(data, 'name', current)
    customer = text_value(data, 'customer_number', current)
    consumption = text_value(data, 'consumption_place', current)
    meter = text_value(data, 'metering_point', current)
    if not name or not customer or not consumption or not meter:
        raise ValueError('Назив, број на корисник, место на потрошувачка и точка на отчитување се задолжителни')
    address = optional(data.get('address')) if 'address' in data else (current['address'] if current else None)
    return {
        'partner_id': int(partner_id),
        'name': name,
        'customer_number': customer,
        'consumption_place': consumption,
        'metering_point': meter,
        'address': address,
    }


@app.route('/api/places', methods=['POST'])
@login_required
def create_place():
    data = request.get_json(silent=True) or {}
    try:
        parsed = read_place(data)
    except (ValueError, TypeError) as exc:
        return bad(str(exc) if str(exc) else 'Невалидно место')
    db = get_db()
    if not db.execute('SELECT id FROM partners WHERE id = ?', (parsed['partner_id'],)).fetchone():
        return bad('Партнерот не постои', 404)
    try:
        cursor = db.execute(
            '''INSERT INTO places
               (partner_id, name, customer_number, consumption_place, metering_point, address)
               VALUES (?, ?, ?, ?, ?, ?)''',
            tuple(parsed.values()),
        )
        db.commit()
    except sqlite3.IntegrityError:
        return bad('Бројот на корисник веќе постои кај тој партнер', 409)
    return jsonify(place_payload(db.execute(
        'SELECT * FROM places WHERE id = ?', (cursor.lastrowid,)
    ).fetchone())), 201


@app.route('/api/places/<int:place_id>', methods=['PUT'])
@login_required
def update_place(place_id):
    data = request.get_json(silent=True) or {}
    db = get_db()
    current = db.execute('SELECT * FROM places WHERE id = ?', (place_id,)).fetchone()
    if not current:
        return bad('Местото не постои', 404)
    try:
        parsed = read_place(data, current)
    except (ValueError, TypeError) as exc:
        return bad(str(exc) if str(exc) else 'Невалидно место')
    try:
        db.execute(
            '''UPDATE places SET partner_id=?, name=?, customer_number=?,
               consumption_place=?, metering_point=?, address=? WHERE id=?''',
            (*parsed.values(), place_id),
        )
        db.commit()
    except sqlite3.IntegrityError:
        return bad('Бројот на корисник веќе постои кај тој партнер', 409)
    return jsonify(place_payload(db.execute(
        'SELECT * FROM places WHERE id = ?', (place_id,)
    ).fetchone()))


@app.route('/api/places/<int:place_id>', methods=['DELETE'])
@login_required
def delete_place(place_id):
    db = get_db()
    used = db.execute(
        'SELECT COUNT(*) FROM invoices WHERE place_id = ?', (place_id,)
    ).fetchone()[0]
    if used:
        return bad('Местото се користи на фактури и не може да се избрише')
    deleted = db.execute('DELETE FROM places WHERE id = ?', (place_id,)).rowcount
    db.commit()
    if not deleted:
        return bad('Местото не постои', 404)
    return jsonify({'ok': True})


def invoice_totals(db, invoice_id, rounding):
    totals = db.execute(
        '''SELECT COALESCE(SUM(amount_without_tax), 0) AS net,
                  COALESCE(SUM(tax_amount), 0) AS tax,
                  COALESCE(SUM(amount_with_tax), 0) AS gross
           FROM invoice_lines WHERE invoice_id = ?''',
        (invoice_id,),
    ).fetchone()
    payable = money(Decimal(str(totals['gross'])) + Decimal(str(rounding or 0)))
    return {
        'amount_without_tax': float(money(totals['net'])),
        'tax_amount': float(money(totals['tax'])),
        'amount_with_tax': float(money(totals['gross'])),
        'rounding': float(money(rounding or 0)),
        'payable': float(payable),
    }


def invoice_header(row, totals):
    item = dict(row)
    item.update(totals)
    item['payment_status_label'] = STATUS_LABELS.get(row['payment_status'], row['payment_status'])
    return item


INVOICE_SELECT = '''
    SELECT i.*, p.name AS partner_name, pl.name AS place_name,
           pl.customer_number, pl.consumption_place
    FROM invoices i
    JOIN partners p ON p.id = i.partner_id
    LEFT JOIN places pl ON pl.id = i.place_id
'''


def load_invoice(db, invoice_id):
    row = db.execute(INVOICE_SELECT + ' WHERE i.id = ?', (invoice_id,)).fetchone()
    if not row:
        return None
    lines = db.execute(
        '''SELECT l.*, t.name AS tax_type_name
           FROM invoice_lines l
           LEFT JOIN tax_types t ON t.id = l.tax_type_id
           WHERE l.invoice_id = ?
           ORDER BY l.line_no, l.id''',
        (invoice_id,),
    ).fetchall()
    item = invoice_header(row, invoice_totals(db, invoice_id, row['rounding']))
    item['lines'] = [dict(line) for line in lines]
    return item


def read_lines(db, lines):
    if not lines:
        raise ValueError('Фактурата треба да има барем една ставка')
    parsed = []
    for index, line in enumerate(lines, start=1):
        article_id = line.get('article_id') or None
        article = None
        if article_id:
            article = db.execute('SELECT * FROM articles WHERE id = ?', (article_id,)).fetchone()
            if not article:
                raise ValueError(f'Ставка {index}: избраната ставка не постои')
        description = (line.get('description') or '').strip() or (article['name'] if article else '')
        if not description:
            raise ValueError(f'Ставка {index}: изберете ставка')
        quantity = parse_number(line.get('quantity'), QTY, label=f'Ставка {index}: количина')
        unit_price = parse_number(line.get('unit_price'), PRICE, label=f'Ставка {index}: цена')
        amount = parse_number(
            line.get('amount_without_tax'), MONEY,
            label=f'Ставка {index}: износ без данок',
        )
        if amount is None:
            if quantity is None or unit_price is None:
                raise ValueError(f'Ставка {index}: внесете износ без данок или количина и цена')
            amount = money(quantity * unit_price)
        tax_id = line.get('tax_type_id')
        if not tax_id:
            raise ValueError(f'Ставка {index}: изберете даночна стапка')
        tax = db.execute('SELECT * FROM tax_types WHERE id = ?', (tax_id,)).fetchone()
        if not tax:
            raise ValueError(f'Ставка {index}: даночната стапка не постои')
        rate = Decimal(str(tax['rate']))
        tax_amount = money(amount * rate)
        parsed.append({
            'line_no': index,
            'article_id': article['id'] if article else None,
            'article_code': article['code'] if article else optional(line.get('article_code')),
            'description': description,
            'quantity': float_or_none(quantity),
            'unit': optional(line.get('unit')) or (article['unit'] if article else None),
            'unit_price': float_or_none(unit_price),
            'amount_without_tax': float(amount),
            'tax_type_id': tax['id'],
            'tax_rate': float(rate),
            'tax_amount': float(tax_amount),
            'amount_with_tax': float(amount + tax_amount),
        })
    return parsed


def apply_rounding(db, invoice_id, rounding):
    if rounding is None:
        gross = invoice_totals(db, invoice_id, 0)['amount_with_tax']
        rounding = float(denar(gross) - money(gross))
    db.execute('UPDATE invoices SET rounding = ? WHERE id = ?', (rounding, invoice_id))


def save_lines(db, invoice_id, lines):
    db.execute('DELETE FROM invoice_lines WHERE invoice_id = ?', (invoice_id,))
    db.executemany(
        '''INSERT INTO invoice_lines
           (invoice_id, line_no, article_id, article_code, description, quantity, unit, unit_price,
            amount_without_tax, tax_type_id, tax_rate, tax_amount, amount_with_tax)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)''',
        [( 
            invoice_id, line['line_no'], line['article_id'], line['article_code'], line['description'],
            line['quantity'], line['unit'], line['unit_price'], line['amount_without_tax'],
            line['tax_type_id'], line['tax_rate'], line['tax_amount'], line['amount_with_tax'],
        ) for line in lines],
    )


def read_invoice(db, data):
    number = (data.get('invoice_number') or '').strip()
    if not number:
        raise ValueError('Бројот на фактура е задолжителен')
    partner_id = data.get('partner_id')
    if not partner_id:
        raise ValueError('Партнерот е задолжителен')
    partner = db.execute('SELECT id FROM partners WHERE id = ?', (partner_id,)).fetchone()
    if not partner:
        raise ValueError('Партнерот не постои')
    place_id = data.get('place_id') or None
    place_count = db.execute(
        'SELECT COUNT(*) FROM places WHERE partner_id = ?', (partner_id,)
    ).fetchone()[0]
    if place_count and not place_id:
        raise ValueError('Изберете место на потрошувачка')
    if place_id:
        place = db.execute(
            'SELECT id FROM places WHERE id = ? AND partner_id = ?',
            (place_id, partner_id),
        ).fetchone()
        if not place:
            raise ValueError('Местото не припаѓа на избраниот партнер')
    status = data.get('payment_status') or 'unpaid'
    if status not in STATUSES:
        raise ValueError('Неважечки статус')
    if data.get('rounding') not in (None, ''):
        rounding = parse_number(data.get('rounding'), MONEY, label='Израмнување')
    else:
        rounding = None
    return {
        'invoice_number': number,
        'issue_date': parse_date(data.get('issue_date'), required=True, label='Датум на издавање'),
        'due_date': parse_date(data.get('due_date'), label='Рок на плаќање'),
        'period_from': parse_date(data.get('period_from'), label='Период од'),
        'period_to': parse_date(data.get('period_to'), label='Период до'),
        'partner_id': int(partner_id),
        'place_id': int(place_id) if place_id else None,
        'payment_ref': optional(data.get('payment_ref')),
        'rounding': None if rounding is None else float(rounding),
        'payment_status': status,
        'notes': optional(data.get('notes')),
    }


INVOICE_COLUMNS = (
    'invoice_number', 'issue_date', 'due_date', 'period_from', 'period_to',
    'partner_id', 'place_id', 'payment_ref', 'rounding', 'payment_status', 'notes',
)


@app.route('/api/invoices', methods=['GET'])
@login_required
def get_invoices():
    clauses = []
    params = []
    if request.args.get('partner_id'):
        clauses.append('i.partner_id = ?')
        params.append(request.args['partner_id'])
    if request.args.get('place_id'):
        clauses.append('i.place_id = ?')
        params.append(request.args['place_id'])
    if request.args.get('status'):
        clauses.append('i.payment_status = ?')
        params.append(request.args['status'])
    if request.args.get('year'):
        clauses.append("strftime('%Y', i.issue_date) = ?")
        params.append(request.args['year'])
    if request.args.get('from'):
        clauses.append('i.issue_date >= ?')
        params.append(request.args['from'])
    if request.args.get('to'):
        clauses.append('i.issue_date <= ?')
        params.append(request.args['to'])
    search = (request.args.get('q') or '').strip()
    if search:
        clauses.append('''
            (i.invoice_number LIKE ? OR p.name LIKE ? OR IFNULL(pl.name, '') LIKE ?
             OR IFNULL(pl.customer_number, '') LIKE ? OR IFNULL(i.notes, '') LIKE ?
             OR IFNULL(i.payment_ref, '') LIKE ?)
        ''')
        like = f'%{search}%'
        params.extend([like] * 6)
    sql = INVOICE_SELECT
    if clauses:
        sql += ' WHERE ' + ' AND '.join(clauses)
    sql += ' ORDER BY i.issue_date DESC, i.invoice_number DESC'
    db = get_db()
    rows = []
    grand = {'amount_without_tax': 0.0, 'tax_amount': 0.0, 'amount_with_tax': 0.0, 'payable': 0.0}
    for row in db.execute(sql, params):
        item = invoice_header(row, invoice_totals(db, row['id'], row['rounding']))
        rows.append(item)
        for key, total in grand.items():
            grand[key] = float(money(Decimal(str(total)) + Decimal(str(item[key]))))
    return jsonify({'invoices': rows, 'totals': grand, 'count': len(rows)})


@app.route('/api/invoices/<int:invoice_id>', methods=['GET'])
@login_required
def get_invoice(invoice_id):
    item = load_invoice(get_db(), invoice_id)
    if not item:
        return bad('Фактурата не постои', 404)
    return jsonify(item)


@app.route('/api/invoices', methods=['POST'])
@login_required
def create_invoice():
    data = request.get_json(silent=True) or {}
    db = get_db()
    try:
        parsed = read_invoice(db, data)
        lines = read_lines(db, data.get('lines') or [])
        cursor = db.execute(
            f'''INSERT INTO invoices ({', '.join(INVOICE_COLUMNS)})
                VALUES ({', '.join('?' for _ in INVOICE_COLUMNS)})''',
            tuple(0 if column == 'rounding' else parsed[column] for column in INVOICE_COLUMNS),
        )
        save_lines(db, cursor.lastrowid, lines)
        apply_rounding(db, cursor.lastrowid, parsed['rounding'])
        db.commit()
    except ValueError as exc:
        db.rollback()
        return bad(str(exc))
    except sqlite3.IntegrityError:
        db.rollback()
        return bad('Фактура со тој број веќе постои кај партнерот', 409)
    return jsonify(load_invoice(db, cursor.lastrowid)), 201


@app.route('/api/invoices/<int:invoice_id>', methods=['PUT'])
@login_required
def update_invoice(invoice_id):
    data = request.get_json(silent=True) or {}
    db = get_db()
    if not db.execute('SELECT id FROM invoices WHERE id = ?', (invoice_id,)).fetchone():
        return bad('Фактурата не постои', 404)
    try:
        parsed = read_invoice(db, data)
        lines = read_lines(db, data.get('lines') or [])
        db.execute(
            f'''UPDATE invoices SET {', '.join(f'{column}=?' for column in INVOICE_COLUMNS)}
                WHERE id=?''',
            (*[0 if column == 'rounding' else parsed[column] for column in INVOICE_COLUMNS], invoice_id),
        )
        save_lines(db, invoice_id, lines)
        apply_rounding(db, invoice_id, parsed['rounding'])
        db.commit()
    except ValueError as exc:
        db.rollback()
        return bad(str(exc))
    except sqlite3.IntegrityError:
        db.rollback()
        return bad('Фактура со тој број веќе постои кај партнерот', 409)
    return jsonify(load_invoice(db, invoice_id))


@app.route('/api/invoices/<int:invoice_id>', methods=['DELETE'])
@login_required
def delete_invoice(invoice_id):
    db = get_db()
    deleted = db.execute('DELETE FROM invoices WHERE id = ?', (invoice_id,)).rowcount
    db.commit()
    if not deleted:
        return bad('Фактурата не постои', 404)
    return jsonify({'ok': True})


def report_filters():
    date_from = request.args.get('from') or None
    date_to = request.args.get('to') or None
    partner_id = request.args.get('partner_id') or None
    return date_from, date_to, partner_id


def date_clause(alias='i'):
    date_from, date_to, partner_id = report_filters()
    clauses, params = [], []
    if date_from:
        clauses.append(f'{alias}.issue_date >= ?')
        params.append(date_from)
    if date_to:
        clauses.append(f'{alias}.issue_date <= ?')
        params.append(date_to)
    if partner_id:
        clauses.append(f'{alias}.partner_id = ?')
        params.append(partner_id)
    return clauses, params, date_from, date_to, partner_id


@app.route('/api/reports/places', methods=['GET'])
@login_required
def report_places():
    db = get_db()
    clauses, params, date_from, date_to, partner_id = date_clause()
    where = (' AND ' + ' AND '.join(clauses)) if clauses else ''
    place_sql = 'SELECT * FROM places'
    place_params = []
    if partner_id:
        place_sql += ' WHERE partner_id = ?'
        place_params.append(partner_id)
    place_sql += ' ORDER BY customer_number'
    rows = []
    grand = {'invoice_count': 0, 'amount_without_tax': 0.0, 'tax_amount': 0.0, 'payable': 0.0}
    for place in db.execute(place_sql, place_params):
        totals = db.execute(
            f'''SELECT COUNT(DISTINCT i.id) AS invoice_count,
                       COALESCE(SUM(l.amount_without_tax), 0) AS net,
                       COALESCE(SUM(l.tax_amount), 0) AS tax
                FROM invoices i
                LEFT JOIN invoice_lines l ON l.invoice_id = i.id
                WHERE i.place_id = ?{where}''',
            (place['id'], *params),
        ).fetchone()
        payable = db.execute(
            f'''SELECT COALESCE(SUM(amount_with_tax), 0) FROM (
                    SELECT i.rounding + COALESCE(SUM(l.amount_with_tax), 0) AS amount_with_tax
                    FROM invoices i
                    LEFT JOIN invoice_lines l ON l.invoice_id = i.id
                    WHERE i.place_id = ?{where}
                    GROUP BY i.id
                )''',
            (place['id'], *params),
        ).fetchone()[0]
        item = {
            'place_id': place['id'],
            'place_name': place['name'],
            'partner_id': place['partner_id'],
            'customer_number': place['customer_number'],
            'consumption_place': place['consumption_place'],
            'invoice_count': totals['invoice_count'],
            'amount_without_tax': float(money(totals['net'])),
            'tax_amount': float(money(totals['tax'])),
            'payable': float(money(payable)),
        }
        rows.append(item)
        grand['invoice_count'] += item['invoice_count']
        for key in ('amount_without_tax', 'tax_amount', 'payable'):
            grand[key] = float(money(Decimal(str(grand[key])) + Decimal(str(item[key]))))
    none_where = 'i.place_id IS NULL' + where
    none_payable = db.execute(
        f'''SELECT COALESCE(SUM(amount_with_tax), 0) FROM (
                SELECT i.rounding + COALESCE(SUM(l.amount_with_tax), 0) AS amount_with_tax
                FROM invoices i
                LEFT JOIN invoice_lines l ON l.invoice_id = i.id
                WHERE {none_where}
                GROUP BY i.id
            )''',
        params,
    ).fetchone()[0]
    none_totals = db.execute(
        f'''SELECT COUNT(DISTINCT i.id) AS invoice_count,
                   COALESCE(SUM(l.amount_without_tax), 0) AS net,
                   COALESCE(SUM(l.tax_amount), 0) AS tax
            FROM invoices i
            LEFT JOIN invoice_lines l ON l.invoice_id = i.id
            WHERE {none_where}''',
        params,
    ).fetchone()
    without_place = {
        'place_id': None,
        'place_name': 'Без место на потрошувачка',
        'customer_number': '',
        'consumption_place': '',
        'invoice_count': none_totals['invoice_count'],
        'amount_without_tax': float(money(none_totals['net'])),
        'tax_amount': float(money(none_totals['tax'])),
        'payable': float(money(none_payable)),
    }
    if without_place['invoice_count']:
        rows.append(without_place)
        grand['invoice_count'] += without_place['invoice_count']
        for key in ('amount_without_tax', 'tax_amount', 'payable'):
            grand[key] = float(money(Decimal(str(grand[key])) + Decimal(str(without_place[key]))))
    return jsonify({
        'from': date_from,
        'to': date_to,
        'partner_id': int(partner_id) if partner_id else None,
        'places': rows,
        'totals': grand,
    })


@app.route('/api/reports/partner', methods=['GET'])
@login_required
def report_partner():
    partner_id = request.args.get('partner_id')
    if not partner_id:
        return bad('Изберете партнер')
    db = get_db()
    partner = db.execute('SELECT * FROM partners WHERE id = ?', (partner_id,)).fetchone()
    if not partner:
        return bad('Партнерот не постои', 404)
    date_from = request.args.get('from') or None
    date_to = request.args.get('to') or None
    clauses = ['i.partner_id = ?']
    params = [partner_id]
    if date_from:
        clauses.append('i.issue_date >= ?')
        params.append(date_from)
    if date_to:
        clauses.append('i.issue_date <= ?')
        params.append(date_to)
    where = ' AND '.join(clauses)
    articles = []
    for row in db.execute(
        f'''SELECT IFNULL(l.article_code, '') AS article_code,
                   l.description,
                   IFNULL(l.unit, '') AS unit,
                   SUM(l.quantity) AS quantity,
                   SUM(CASE WHEN l.quantity IS NULL THEN 1 ELSE 0 END) AS missing_quantity,
                   SUM(l.amount_without_tax) AS amount_without_tax,
                   SUM(l.tax_amount) AS tax_amount,
                   SUM(l.amount_with_tax) AS amount_with_tax,
                   COUNT(*) AS line_count
            FROM invoice_lines l
            JOIN invoices i ON i.id = l.invoice_id
            WHERE {where}
            GROUP BY IFNULL(l.article_code, ''), l.description, IFNULL(l.unit, '')
            ORDER BY amount_with_tax DESC, l.description''',
        params,
    ):
        item = dict(row)
        if row['quantity'] is None:
            item['quantity'] = None
        else:
            item['quantity'] = float(Decimal(str(row['quantity'])).quantize(QTY, rounding=ROUND_HALF_UP))
            if row['missing_quantity'] and item['quantity'] == 0:
                item['quantity'] = None
        articles.append(item)
    invoices = []
    for row in db.execute(
        f'''SELECT i.*, pl.name AS place_name, pl.customer_number
            FROM invoices i
            LEFT JOIN places pl ON pl.id = i.place_id
            WHERE {where}
            ORDER BY i.issue_date, i.invoice_number''',
        params,
    ):
        invoices.append(invoice_header(row, invoice_totals(db, row['id'], row['rounding'])))
    quantity_total = None
    units = {article['unit'] for article in articles if article['quantity'] is not None}
    if len(units) == 1 and articles:
        quantity_total = float(sum(Decimal(str(article['quantity'] or 0)) for article in articles))
    totals = {
        'invoice_count': len(invoices),
        'quantity': quantity_total,
        'unit': next(iter(units)) if len(units) == 1 else None,
        'amount_without_tax': float(sum((Decimal(str(item['amount_without_tax'])) for item in invoices), Decimal(0))),
        'tax_amount': float(sum((Decimal(str(item['tax_amount'])) for item in invoices), Decimal(0))),
        'rounding': float(sum((Decimal(str(item['rounding'])) for item in invoices), Decimal(0))),
        'amount_with_tax': float(sum((Decimal(str(item['amount_with_tax'])) for item in invoices), Decimal(0))),
        'payable': float(sum((Decimal(str(item['payable'])) for item in invoices), Decimal(0))),
    }
    for key in ('amount_without_tax', 'tax_amount', 'rounding', 'amount_with_tax', 'payable'):
        totals[key] = float(money(totals[key]))
    return jsonify({
        'partner': dict(partner),
        'from': date_from,
        'to': date_to,
        'articles': articles,
        'invoices': invoices,
        'totals': totals,
    })


def workbook(title, headers, rows):
    book = Workbook()
    sheet = book.active or book.create_sheet()
    sheet.title = title[:31]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for row in rows:
        sheet.append(row)
    for index, column in enumerate(sheet.iter_cols(min_row=1, max_row=sheet.max_row, max_col=sheet.max_column), start=1):
        letter = get_column_letter(index)
        width = max(len(str(cell.value or '')) for cell in column)
        sheet.column_dimensions[letter].width = min(max(width + 2, 12), 42)
    output = io.BytesIO()
    book.save(output)
    output.seek(0)
    return output


PARTNER_FIELDS = [
    ('name', 'Назив'),
    ('edb', 'ЕДБ'),
    ('embs', 'ЕМБС'),
    ('address', 'Адреса'),
    ('city', 'Град'),
    ('country', 'Држава'),
    ('phone', 'Телефон'),
    ('mobile', 'Мобилен'),
    ('email', 'Е-пошта'),
    ('notes', 'Забелешка'),
]
FONT_CANDIDATES = (
    r'C:\Windows\Fonts\arial.ttf',
    r'C:\Windows\Fonts\segoeui.ttf',
    '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
)


def partner_cards(db, partner_id=None):
    sql = 'SELECT * FROM partners'
    params = []
    if partner_id:
        sql += ' WHERE id = ?'
        params.append(partner_id)
    sql += ' ORDER BY name'
    cards = []
    for partner in db.execute(sql, params):
        places = db.execute(
            '''SELECT name, customer_number, consumption_place, metering_point, address
               FROM places WHERE partner_id = ? ORDER BY customer_number''',
            (partner['id'],),
        ).fetchall()
        cards.append((dict(partner), [dict(place) for place in places], partner_accounts(db, partner['id'])))
    return cards


def partner_pairs(partner, places, accounts=None):
    pairs = [(label, partner.get(key) or '') for key, label in PARTNER_FIELDS if key != 'name']
    if not accounts:
        pairs.append(('Банкарска сметка', ''))
    for index, account in enumerate(accounts or [], start=1):
        bank = account.get('bank') or ''
        number = account.get('account_number') or ''
        pairs.append((f'Банкарска сметка {index}', f'{bank} {number}'.strip()))
    if not places:
        pairs.append(('Места на потрошувачка', 'Нема места'))
        return pairs
    for index, place in enumerate(places, start=1):
        pairs.extend([
            (f'Место {index}', place['name'] or ''),
            ('Број на корисник', place['customer_number'] or ''),
            ('Место на потрошувачка', place['consumption_place'] or ''),
            ('Точка на отчитување', place['metering_point'] or ''),
            ('Адреса на место', place['address'] or ''),
        ])
    return pairs


def partner_excel(cards):
    book = Workbook()
    sheet = book.active or book.create_sheet()
    sheet.title = 'Партнери'
    for partner, places, accounts in cards:
        sheet.append([partner['name'], ''])
        for cell in sheet[sheet.max_row]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
        sheet.merge_cells(start_row=sheet.max_row, start_column=1, end_row=sheet.max_row, end_column=2)
        for label, value in partner_pairs(partner, places, accounts):
            sheet.append([label, value])
        sheet.append(['', ''])
    sheet.column_dimensions['A'].width = 28
    sheet.column_dimensions['B'].width = 62
    output = io.BytesIO()
    book.save(output)
    output.seek(0)
    return output


def find_font():
    for path in FONT_CANDIDATES:
        if os.path.exists(path):
            return path
    raise FileNotFoundError('Нема фонт за македонски текст во PDF')


def wrap_text(text, font, size, width):
    lines = []
    for paragraph in str(text or '').split('\n') or ['']:
        line = ''
        for word in paragraph.split(' ') or ['']:
            trial = word if not line else line + ' ' + word
            if font.text_length(trial, size) <= width or not line:
                line = trial
            else:
                lines.append(line)
                line = word
        lines.append(line)
    return lines or ['']


def partner_pdf(cards):
    document = fitz.open()
    font = fitz.Font(fontfile=find_font())
    widths = [160, 360]
    size = 10
    ink = (0.11, 0.16, 0.14)
    page = document.new_page(width=595, height=842)
    page.insert_font(fontname='mk', fontfile=find_font())
    y = 36

    def new_page():
        nonlocal page, y
        page = document.new_page(width=595, height=842)
        page.insert_font(fontname='mk', fontfile=find_font())
        y = 36

    def draw_row(values, header=False):
        nonlocal y
        font_size = 12 if header else size
        wrapped = [wrap_text(value, font, font_size, width - 10) for value, width in zip(values, widths)]
        row_height = max(len(lines) for lines in wrapped) * (font_size + 3) + 8
        if y + row_height > 810:
            new_page()
        x = 36
        fill = (0.12, 0.30, 0.23) if header else (0.93, 0.91, 0.86)
        text_color = (1, 1, 1) if header else ink
        for lines, width in zip(wrapped, widths):
            page.draw_rect(fitz.Rect(x, y, x + width, y + row_height), color=(0.72, 0.68, 0.58), fill=fill, width=0.4)
            text_y = y + font_size + 3
            for line in lines:
                if line:
                    page.insert_text((x + 5, text_y), line, fontname='mk', fontsize=font_size, color=text_color)
                text_y += font_size + 3
            x += width
        y += row_height

    new_page()
    first = True
    for partner, places, accounts in cards:
        if not first:
            y += 16
        first = False
        draw_row([partner['name'], ''], True)
        for label, value in partner_pairs(partner, places, accounts):
            draw_row([label, value])
    output = io.BytesIO()
    document.save(output)
    document.close()
    output.seek(0)
    return output


def excel_response(output, filename):
    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


@app.route('/api/export/invoices')
@login_required
def export_invoices():
    data = get_invoices().get_json()
    rows = []
    for item in data['invoices']:
        rows.append([
            item['invoice_number'],
            display_date(item['issue_date']),
            display_date(item['due_date']),
            item['partner_name'],
            item['place_name'] or '',
            item['customer_number'] or '',
            item['amount_without_tax'],
            item['tax_amount'],
            item['rounding'],
            item['payable'],
            item['payment_status_label'],
        ])
    rows.append(['', '', '', '', '', 'Вкупно', data['totals']['amount_without_tax'],
                 data['totals']['tax_amount'], '', data['totals']['payable'], ''])
    output = workbook('Фактури', [
        'Број на фактура', 'Датум', 'Рок', 'Партнер', 'Место', 'Број на корисник',
        'Без данок', 'Данок', 'Израмнување', 'За плаќање', 'Статус',
    ], rows)
    return excel_response(output, 'fakturi.xlsx')


@app.route('/api/export/places')
@login_required
def export_places():
    data = report_places().get_json()
    rows = [[
        item['place_name'], item['customer_number'], item['consumption_place'],
        item['invoice_count'], item['amount_without_tax'], item['tax_amount'], item['payable'],
    ] for item in data['places']]
    totals = data['totals']
    rows.append(['Вкупно', '', '', totals['invoice_count'], totals['amount_without_tax'],
                 totals['tax_amount'], totals['payable']])
    output = workbook('По места', [
        'Место', 'Број на корисник', 'Место на потрошувачка', 'Фактури',
        'Без данок', 'Данок', 'За плаќање',
    ], rows)
    return excel_response(output, 'izvestaj-mesta.xlsx')


@app.route('/api/export/partner')
@login_required
def export_partner():
    response = report_partner()
    if isinstance(response, tuple) or getattr(response, 'status_code', 200) >= 400:
        return response
    data = response.get_json()
    rows = [[
        article['article_code'], article['description'], article['unit'],
        article['quantity'] if article['quantity'] is not None else '',
        article['amount_without_tax'], article['tax_amount'], article['amount_with_tax'],
    ] for article in data['articles']]
    totals = data['totals']
    rows.append(['', 'Вкупно ставки', totals['unit'] or '',
                 totals['quantity'] if totals['quantity'] is not None else '',
                 totals['amount_without_tax'], totals['tax_amount'], totals['amount_with_tax']])
    rows.append(['', 'Израмнување', '', '', '', '', totals['rounding']])
    rows.append(['', 'За плаќање', '', '', '', '', totals['payable']])
    output = workbook('По партнер', [
        'Шифра', 'Назив', 'ЕМ', 'Количина', 'Без данок', 'Данок', 'Со данок',
    ], rows)
    return excel_response(output, 'izvestaj-partner.xlsx')


@app.route('/api/export/partner-data.<kind>')
@login_required
def export_partner_data(kind):
    if kind not in ('xlsx', 'pdf'):
        return bad('Непознат формат', 404)
    db = get_db()
    partner_id = request.args.get('partner_id')
    cards = partner_cards(db, partner_id)
    if partner_id and not cards:
        return bad('Партнерот не постои', 404)
    filename = 'partner.' + kind if partner_id else 'partneri.' + kind
    if kind == 'xlsx':
        return excel_response(partner_excel(cards), filename)
    try:
        output = partner_pdf(cards)
    except FileNotFoundError as exc:
        return bad(str(exc))
    return send_file(output, as_attachment=True, download_name=filename, mimetype='application/pdf')


if __name__ == '__main__':
    init_db()
    app._db_ready = True
    print('Влезни фактури: http://localhost:5050')
    app.run(host='127.0.0.1', port=5050, debug=False)
